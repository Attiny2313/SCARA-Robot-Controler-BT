import re
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BLUE = '1F4E79'
BLUE2 = '2F5597'
DARK = '333333'
GRAY = '6B7280'
BORDER = 'B7C9DC'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    trPr.append(header)


def set_table_borders(table, color=BORDER, size=6):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = borders.find(qn(f'w:{edge}'))
        if el is None:
            el = OxmlElement(f'w:{edge}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:color'), color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def set_font(run, name='Arial', size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text, base_size=10.5):
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`)')
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, size=base_size, color=DARK)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, bold=True, color=DARK)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name='Courier New', size=9.5, color='1F2937')
            rPr = run._r.get_or_add_rPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'F1F5F9')
            rPr.append(shd)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=base_size, color=DARK)


def configure_document(doc, lang):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(18)
    sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(20)
    sec.right_margin = Mm(20)
    sec.header_distance = Mm(8)
    sec.footer_distance = Mm(8)

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    specs = {
        'Heading 1': (15, BLUE, 14, 6),
        'Heading 2': (12, BLUE2, 10, 4),
        'Heading 3': (10.5, '44546A', 8, 3),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = 'Arial'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('SCARA Robot Controller — ' + ('dokumentacja techniczna' if lang == 'PL' else 'технічна документація'))
    set_font(run, size=8, color='808080')

    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Attiny2313/SCARA-Robot-Controler-BT   •   ')
    set_font(run, size=8, color='808080')
    add_page_number(p)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith('|'):
        rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in rows[1]):
        rows.pop(1)
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [Mm(105), Mm(55)] if cols == 2 else [Mm(160 / cols)] * cols
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.width = widths[min(ci, len(widths) - 1)]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, row[ci] if ci < len(row) else '', base_size=9.2)
            for run in p.runs:
                if ri == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if ri == 0:
                set_cell_shading(cell, BLUE)
            elif ri % 2 == 0:
                set_cell_shading(cell, 'F7FAFC')
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc, lines, lang):
    titles = []
    idx = 0
    while idx < len(lines) and idx < 3 and lines[idx].startswith('#'):
        titles.append(lines[idx].lstrip('#').strip())
        idx += 1

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    run = p.add_run(titles[0] if titles else '')
    set_font(run, size=25, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(titles[1] if len(titles) > 1 else '')
    set_font(run, size=18, bold=True, color=BLUE2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(titles[2] if len(titles) > 2 else '')
    set_font(run, size=13, bold=True, color='44546A')
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    while idx < len(lines):
        text = lines[idx].rstrip()
        if not text:
            idx += 1
            continue
        if text.startswith('>') or text.strip() == '---':
            break
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(20)
        p.paragraph_format.right_indent = Mm(20)
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, text, base_size=10.5)
        idx += 1

    while idx < len(lines) and not lines[idx].strip():
        idx += 1
    if idx < len(lines) and lines[idx].lstrip().startswith('>'):
        note = []
        while idx < len(lines) and lines[idx].lstrip().startswith('>'):
            note.append(lines[idx].lstrip()[1:].strip())
            idx += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(18)
        p.paragraph_format.right_indent = Mm(18)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(' '.join(note))
        set_font(run, size=9.5, italic=True, color=GRAY)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '8')
        left.set(qn('w:color'), BLUE2)
        pbdr.append(left)
        pPr.append(pbdr)

    doc.add_page_break()
    return idx


def add_contents(doc, lines, lang):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(0)
    p.add_run('SPIS TREŚCI' if lang == 'PL' else 'ЗМІСТ')
    for line in lines:
        if line.startswith('# ') and not line.startswith('##'):
            text = line[2:].strip()
            if re.match(r'^\d+\.', text):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(6)
                p.paragraph_format.first_line_indent = Mm(-4)
                p.paragraph_format.space_after = Pt(2.5)
                run = p.add_run(text)
                set_font(run, size=9.5, color='374151')
    doc.add_page_break()


def convert(md_path, out_path, lang):
    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    doc = Document()
    configure_document(doc, lang)
    start = add_cover(doc, lines, lang)
    add_contents(doc, lines, lang)

    i = start
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == '---':
            i += 1
            continue

        if i < 3 and s.startswith('#'):
            i += 1
            continue

        heading = re.match(r'^(#{1,3})\s+(.*)$', s)
        if heading:
            level = len(heading.group(1))
            p = doc.add_paragraph(style=f'Heading {level}')
            add_inline(p, heading.group(2).strip(), base_size={1: 15, 2: 12, 3: 10.5}[level])
            i += 1
            continue

        if s.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if s.startswith('>'):
            parts = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                parts.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(8)
            p.paragraph_format.right_indent = Mm(4)
            run = p.add_run(' '.join(parts))
            set_font(run, size=9.7, italic=True, color=GRAY)
            continue

        if re.match(r'^[-*]\s+', s):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(1.5)
            add_inline(p, re.sub(r'^[-*]\s+', '', s), base_size=10.3)
            i += 1
            continue

        ordered = re.match(r'^(\d+)\.\s+(.*)$', s)
        if ordered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(7)
            p.paragraph_format.first_line_indent = Mm(-5)
            p.paragraph_format.space_after = Pt(1.5)
            run = p.add_run(f'{ordered.group(1)}.  ')
            set_font(run, size=10.3)
            add_inline(p, ordered.group(2), base_size=10.3)
            i += 1
            continue

        parts = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if (nxt == '---' or nxt.startswith('#') or nxt.startswith('|') or nxt.startswith('>')
                    or re.match(r'^[-*]\s+', nxt) or re.match(r'^\d+\.\s+', nxt)):
                break
            parts.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        add_inline(p, ' '.join(part.rstrip() for part in parts), base_size=10.5)

    doc.save(out_path)


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    docs = root / 'docs'
    convert(docs / 'DOKUMENTACJA_TECHNICZNA_PL.md', docs / 'DOKUMENTACJA_TECHNICZNA_PL.docx', 'PL')
    convert(docs / 'DOKUMENTACJA_TECHNICZNA_UA.md', docs / 'DOKUMENTACJA_TECHNICZNA_UA.docx', 'UA')
    print('Generated technical DOCX files.')
